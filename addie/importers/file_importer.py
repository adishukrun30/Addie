"""Import bank / credit-card exports (CSV or XLSX), including Hebrew headers.

Israeli exports usually have a few preamble rows before the header, use
dd/mm/yyyy dates, and either one signed amount column or separate
debit (חובה) / credit (זכות) columns.
"""
import csv
import io
import re
from dataclasses import asdict, dataclass
from datetime import date, datetime

from . import ParsedTransaction

HEADER_ALIASES = {
    "date": ["תאריך עסקה", "תאריך רכישה", "תאריך", "date", "transaction date"],
    "description": ["שם בית העסק", "שם בית עסק", "תיאור", "פרטים", "הפעולה", "תיאור הפעולה", "description", "merchant"],
    "amount": ["סכום חיוב", "סכום העסקה", "סכום עסקה", "סכום", "amount"],
    "debit": ["חובה", "debit"],
    "credit": ["זכות", "credit"],
}

DATE_FORMATS = ["%d/%m/%Y", "%d/%m/%y", "%d.%m.%Y", "%d.%m.%y", "%d-%m-%Y", "%Y-%m-%d", "%d/%m/%Y %H:%M"]
MAX_HEADER_SCAN = 40


class ImportError_(ValueError):
    pass


@dataclass
class Mapping:
    date: int | None = None
    description: int | None = None
    amount: int | None = None
    debit: int | None = None
    credit: int | None = None
    header_row: int = 0
    invert: bool = False  # True for card exports where a positive amount is a charge

    def is_complete(self) -> bool:
        has_amount = self.amount is not None or self.debit is not None or self.credit is not None
        return self.date is not None and self.description is not None and has_amount

    def to_dict(self) -> dict:
        return asdict(self)


def read_rows(filename: str, content: bytes) -> list[list[str]]:
    if filename.lower().endswith((".xlsx", ".xlsm")):
        return _read_xlsx(content)
    return _read_csv(content)


def _read_csv(content: bytes) -> list[list[str]]:
    for encoding in ("utf-8-sig", "cp1255"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ImportError_("לא הצלחתי לזהות את הקידוד של הקובץ")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return [[c.strip() for c in row] for row in csv.reader(io.StringIO(text), dialect)]


def _read_xlsx(content: bytes) -> list[list[str]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    rows = []
    for row in ws.iter_rows(values_only=True):
        rows.append(["" if v is None else (v.strftime("%Y-%m-%d") if isinstance(v, (datetime, date)) else str(v)).strip() for v in row])
    wb.close()
    return rows


def _normalize_header(value: str) -> str:
    return re.sub(r"\s+", " ", value.replace("‏", "").replace("‎", "")).strip().lower()


def detect_mapping(rows: list[list[str]]) -> Mapping:
    best = Mapping()
    best_score = 0
    for idx, row in enumerate(rows[:MAX_HEADER_SCAN]):
        headers = [_normalize_header(c) for c in row]
        mapping = Mapping(header_row=idx)
        score = 0
        for field, aliases in HEADER_ALIASES.items():
            col = _find_column(headers, aliases)
            if col is not None:
                setattr(mapping, field, col)
                score += 1
        if mapping.is_complete() and score > best_score:
            best, best_score = mapping, score
    return best


def _find_column(headers: list[str], aliases: list[str]) -> int | None:
    # exact match first, then "header starts with alias" (e.g. "סכום חיוב ₪")
    for alias in aliases:
        if alias in headers:
            return headers.index(alias)
    for alias in aliases:
        for i, h in enumerate(headers):
            if h.startswith(alias):
                return i
    return None


def parse_date(value: str) -> str | None:
    value = value.strip()
    if not value:
        return None
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(value, fmt).date().isoformat()
        except ValueError:
            continue
    try:
        return datetime.fromisoformat(value).date().isoformat()
    except ValueError:
        return None


def parse_amount(value: str) -> int | None:
    """'1,234.50' -> 123450 agorot. Handles ₪, trailing minus and (parentheses)."""
    value = value.replace("₪", "").replace("NIS", "").replace(",", "").replace("‏", "").replace("‎", "").strip()
    if not value:
        return None
    negative = False
    if value.startswith("(") and value.endswith(")"):
        negative, value = True, value[1:-1]
    if value.endswith("-"):
        negative, value = True, value[:-1]
    if value.startswith("-"):
        negative, value = not negative, value[1:]
    try:
        cents = round(float(value) * 100)
    except ValueError:
        return None
    return -cents if negative else cents


def parse(rows: list[list[str]], mapping: Mapping) -> tuple[list[ParsedTransaction], int]:
    """Returns (transactions, skipped_row_count). Rows without a valid date/amount are skipped
    (preamble, totals, blank lines)."""
    if not mapping.is_complete():
        raise ImportError_("חסרות עמודות: צריך תאריך, תיאור וסכום")
    out, skipped = [], 0

    def cell(row, col):
        return row[col] if col is not None and col < len(row) else ""

    for row in rows[mapping.header_row + 1:]:
        if not any(row):
            continue
        tx_date = parse_date(cell(row, mapping.date))
        description = cell(row, mapping.description)
        if mapping.amount is not None:
            amount = parse_amount(cell(row, mapping.amount))
        else:
            debit = parse_amount(cell(row, mapping.debit)) or 0
            credit = parse_amount(cell(row, mapping.credit)) or 0
            amount = credit - abs(debit) if (debit or credit) else None
        if tx_date is None or amount is None or not description:
            skipped += 1
            continue
        if mapping.invert:
            amount = -amount
        out.append(ParsedTransaction(date=tx_date, description=description, amount=amount))
    return out, skipped
